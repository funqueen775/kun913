from docx import Document
from docx.oxml.ns import qn


PATH = r"C:\Users\23342\Desktop\project\ai_town\workplace_town3\docs\职场小镇开发文档V1.docx"


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run(value)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


document = Document(PATH)
timeline = document.tables[1]
updates = [
    ("第 1 月初 白天", "B 科技丘 你的工位", "新员工手册"),
    ("第 3 月初 白天", "A 总部 评审室", "技术选型"),
    ("第 5 月初 白天", "B 科技丘 工位", "评测集要不要建"),
    ("第 7 月初 深夜", "B 科技丘 工位", "AI 代码标注"),
]

for row_index, (date_text, location, event_title) in enumerate(updates, start=1):
    row = timeline.rows[row_index]
    set_cell_text(row.cells[1], date_text)
    set_cell_text(row.cells[3], location)
    set_cell_text(row.cells[4], event_title)

document.save(PATH)
